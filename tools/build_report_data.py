"""결과지 FINAL pptx 8종 → js/report-data.js

자료/설문 결과지 FINAL/설문 결과지 FINAL/*.pptx 의 슬라이드 문구를 그대로 읽어
설문 결과 화면(js/report.js)이 쓰는 데이터로 만든다. 원고가 바뀌면 이 스크립트만 다시 돌린다.

  python tools/build_report_data.py
"""
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT.parent / "자료" / "설문 결과지 FINAL" / "설문 결과지 FINAL"
OUT = ROOT / "js" / "report-data.js"

FILES = {
    "early": "결과지_초기노화형_FINAL.pptx",
    "tired": "결과지_피로형_FINAL.pptx",
    "fine": "결과지_잔주름형_FINAL.pptx",
    "sag": "결과지_변형처짐형_FINAL.pptx",
    "muscle": "결과지_근육형_FINAL.pptx",
    "mixed_fine": "결과지_혼합형_예시A_눈가잔주름_턱볼처짐_FINAL.pptx",
    "mixed_tired": "결과지_혼합형_예시B_눈밑피로_턱볼처짐_FINAL.pptx",
    "mixed_muscle": "결과지_혼합형_예시C_미간표정_턱볼처짐_FINAL.pptx",
}


def slides(path):
    """슬라이드마다 텍스트 상자 목록(각 상자는 줄 목록)."""
    out = []
    for sl in Presentation(path).slides:
        boxes = []
        for sh in sl.shapes:
            if sh.has_text_frame:
                lines = [p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip()]
                if lines:
                    boxes.append(lines)
        out.append(boxes)
    return out


def body(boxes):
    """'로그인 해야…' 와 제목 상자를 뺀 본문 상자 중 가장 긴 것."""
    cand = [b for b in boxes if not b[0].startswith("로그인 해야") and len(b) > 1]
    return max(cand, key=lambda b: sum(len(x) for x in b)) if cand else []


def clean(s):
    s = re.sub(r"[-]", "", s)   # pptx 기호 글꼴(Wingdings 등)의 사용자 정의 문자 — 화면에선 네모로 깨진다
    return re.sub(r"\s+", " ", s).strip()


def parse(path):
    S = slides(path)
    d = {}

    # 3 — 진단 이유
    b = body(S[2])
    i = next(k for k, x in enumerate(b) if x.startswith("고객님의 설문에서도"))
    d["reason"] = [clean(x) for x in b[:i]]
    ev = [x for x in b[i + 1:] if x.startswith("✓")]
    d["evidence"] = [clean(x.lstrip("✓")) for x in ev]
    rest = [x for x in b[i + 1:] if not x.startswith("✓")]
    note = clean(" ".join(rest))
    d["reasonNote"] = re.sub(r"^답변이 많았습니다\.\s*", "", note)

    # 4 — 집중 케어 부위 TOP3
    b = body(S[3])
    areas, expl, intro = [], [], []
    for x in b:
        m = re.match(r"^[①②③]\s*([A-Z]\d+)\s+(.+?)(?:\s*\((.+?)\))?\s*$", x)
        if m:
            areas.append({"code": m.group(1), "name": clean(m.group(2)), "pattern": m.group(3) or ""})
        elif areas:
            expl.append(clean(x))
        else:
            intro.append(clean(x))
    d["areaIntro"] = " ".join(intro)
    d["areas"] = areas
    d["areaText"] = expl

    # 5 — 홈케어 루틴 + 성분
    b = body(S[4])
    t = next(k for k, x in enumerate(b) if "Tip" in x)
    d["homecare"] = [clean(x) for x in b[:t]]
    d["ingredients"] = [clean(x) for x in b[t + 2:]]

    # 6 — 생활습관 + 디바이스 추천
    life = next(bx for bx in S[5] if bx[0].startswith("생활 습관"))
    d["lifestyle"] = [clean(x) for x in life[1:]]
    devrec = [bx for bx in S[5] if len(bx) == 1 and len(bx[0]) > 80]
    d["deviceRec"] = clean(devrec[0][0]) if devrec else ""

    # 7 — 고주파(RF) 기술과 케어클 (끝의 공통 안내 두 문단은 뺀다)
    b = body(S[6])
    d["rf"] = [clean(x) for x in b if not x.startswith(("뷰티 디바이스 별로", "만약 고객님이"))]

    # 8 — 맞춤 사용법 / 관리 순서 / 권장 빈도
    b = body(S[7])
    o = next(k for k, x in enumerate(b) if x.startswith("<관리 순서>"))
    d["usage"] = [clean(x) for x in b[:o] if not x.startswith("케어클 TECH FIT은")]
    d["order"] = [clean(x) for x in b[o + 1].split("→")]
    f = next(k for k, x in enumerate(b) if x.startswith("<권장 빈도>"))
    d["frequency"] = clean(b[f + 1].replace("·", "·"))

    # 9 — SHOT 가이드 + 마무리 모드
    b = body(S[8])
    shots, modes, total = [], [], None
    for k, x in enumerate(b):
        m = re.match(r"^([A-Z]\d+)\s+(.+?)\s*:\s*(.+?)\s*·\s*(.+?)\s*·\s*(\d+)샷", x)
        if m:
            shots.append({"code": m.group(1), "name": clean(m.group(2)), "mode": clean(m.group(3)),
                          "dir": clean(m.group(4)), "count": int(m.group(5))})
            continue
        m = re.match(r"^총\s*SHOT\s*(\d+)\s*샷", x)
        if m:
            total = int(m.group(1))
            continue
        m = re.match(r"^\d단계\s*:\s*((?!SHOT).+?모드)\s*(?:\((필수|선택)\))?\s*$", x)
        if m and k + 1 < len(b):
            nx = clean(b[k + 1])
            dur = re.match(r"^권장\s*([\d~]+분)\s*·\s*(.+)$", nx)
            # 원고에 필수/선택 표기가 없으면 배지를 달지 않는다(None)
            req = {"필수": True, "선택": False}.get(m.group(2))
            modes.append({"name": clean(m.group(1)), "required": req,
                          "duration": dur.group(1) if dur else "", "desc": dur.group(2) if dur else nx})
    d["shots"], d["shotTotal"], d["modes"] = shots, total, modes

    # 11 — 바쁜 날
    b = body(S[10])
    d["busy"] = [clean(x) for x in b[1:]]

    # 12 — 후기
    b = [clean(x) for x in sum(S[11], [])]
    revs, cur = [], None
    for x in b:
        if x.startswith("★"):
            cur = {"quote": "", "type": "", "concern": "", "period": ""}
            revs.append(cur)
        elif cur is not None:
            if x.startswith("“") and not cur["quote"]:
                cur["quote"] = x.strip("“”\"")
            elif x.startswith("—"):
                cur["type"] = x.lstrip("— ").strip()
            elif "주요 고민" in x:
                cur["concern"] = x.split(":", 1)[1].strip()
            elif "관리 약" in x:
                cur["period"] = x.lstrip("· ").strip()
            elif x.startswith("〔"):
                cur = None
    d["reviews"] = [r for r in revs if r["quote"]][:2]
    return d


def main():
    data = {k: parse(SRC / f) for k, f in FILES.items()}
    js = ("/* 자동 생성 — tools/build_report_data.py (원본: 자료/설문 결과지 FINAL/*.pptx). 직접 고치지 말 것. */\n"
          "window.ckuReportData = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n")
    OUT.write_text(js, encoding="utf-8")
    for k, v in data.items():
        print(k, "areas", len(v["areas"]), "shots", len(v["shots"]), "modes", len(v["modes"]),
              "reviews", len(v["reviews"]), "ingr", len(v["ingredients"]), "life", len(v["lifestyle"]))


if __name__ == "__main__":
    sys.exit(main())
